"""Seed students from '5AM ROT 1 2026-2027.xlsx' into level 4 / rot1.

Idempotent: a student already present with the same (nom, prenom, level,
rotation) is skipped, so re-runs are safe. Run from the project root:
    python scripts/seed_students_rot1.py
"""
import os
import psycopg
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LEVEL, ROTATION = "4", "rot1"
XLSX = os.path.join(os.path.dirname(ROOT), "5AM ROT 1 2026-2027.xlsx")


def env(name):
    with open(os.path.join(ROOT, ".env.local"), encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line.startswith(name + "="):
                return line.split("=", 1)[1].strip()
    return None


PASSWORD = env("SUPABASE_DB_PASSWORD")
REF = "mnsfehjctnjprqsaqart"
DSN = f"postgresql://postgres.{REF}:{PASSWORD}@aws-0-eu-west-2.pooler.supabase.com:5432/postgres"


def read_xlsx():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[".ROT1"]
    rows = []
    for r in range(2, ws.max_row + 1):
        nom = (ws.cell(row=r, column=3).value or "").strip()
        prenom = (ws.cell(row=r, column=4).value or "").strip()
        if nom and prenom:
            rows.append((nom.upper(), prenom))
    return rows


def main():
    students = read_xlsx()
    print(f"xlsx rows: {len(students)}")

    with psycopg.connect(DSN, connect_timeout=15, sslmode="require") as conn:
        with conn.cursor() as cur:
            inserted = 0
            skipped = 0
            for nom, prenom in students:
                cur.execute(
                    """
                    insert into public.students (nom, prenom, level, rotation)
                    select %s, %s, %s, %s
                    where not exists (
                      select 1 from public.students
                      where nom = %s and prenom = %s and level = %s and rotation = %s
                    )
                    """,
                    (nom, prenom, LEVEL, ROTATION, nom, prenom, LEVEL, ROTATION),
                )
                if cur.rowcount == 1:
                    inserted += 1
                else:
                    skipped += 1
            print(f"inserted: {inserted}  skipped (already present): {skipped}")

            # Readback = verification
            cur.execute(
                """
                select nom, prenom from public.students
                where level = %s and rotation = %s
                order by nom, prenom
                """,
                (LEVEL, ROTATION),
            )
            rows = cur.fetchall()
            print(f"level {LEVEL}/rot{ROTATION} students now in db: {len(rows)}")
            for nom, prenom in rows[:8]:
                print(f"  {nom} {prenom}")
            if len(rows) > 8:
                print(f"  … and {len(rows) - 8} more")
            cur.execute(
                "select level, rotation, count(*) from public.students group by level, rotation order by level, rotation"
            )
            print("\nby level/rotation:")
            for r in cur.fetchall():
                print(f"  {r[0]}/{r[1]}: {r[2]}")
            cur.execute("select count(*) from public.students where joined")
            print(f"\njoined so far: {cur.fetchone()[0]}")


if __name__ == "__main__":
    main()