"""One-shot: import ROT2 group + students from the xlsx into Supabase.

Reads .env.local for SUPABASE_URL + SUPABASE_SERVICE_ROLE_KEY (never prints
secrets), parses the ROT2 sheet, inserts the missing group row and student
rows, then verifies by reading back. Idempotent: skips already-present rows.
"""
import json
import re
import sys
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT.parent / "5AM ROT 2 2026-2027.xlsx"
LABEL = "M22/25 Rotation 2"
CHAT_ID = "-1004380365230"
LEVEL, ROTATION = "5", "rot2"


def load_env(path: Path) -> dict:
    env = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            m = re.match(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)$", line)
            if m:
                env[m[1]] = m[2].strip().strip('"').strip("'")
    return env


env = load_env(ROOT / ".env.local")
url, key = env.get("SUPABASE_URL"), env.get("SUPABASE_SERVICE_ROLE_KEY")
if not url or not key:
    sys.exit("missing SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY in .env.local")


def rest(method: str, path: str, body=None, params=""):
    req = urllib.request.Request(
        f"{url}/rest/v1/{path}{params}",
        data=json.dumps(body).encode() if body is not None else None,
        method=method,
        headers={
            "apikey": key,
            "Authorization": f"Bearer {key}",
            "Content-Type": "application/json",
        },
    )
    try:
        with urllib.request.urlopen(req) as resp:
            raw = resp.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as e:
        return {"__http_error__": e.code, "__body__": e.read().decode(errors="replace")}


# --- 1) parse xlsx ---------------------------------------------------------
wb = load_workbook(XLSX, data_only=True)
ws = wb["ROT2"]
rows = list(ws.iter_rows(values_only=True))
assert rows and rows[0][:4] == ("N°", "Matricule", "Nom", "Prénom"), rows[0]
students = []
for r in rows[1:]:
    if r[0] is None:
        continue
    nom = str(r[2]).strip() if r[2] is not None else ""
    prenom = str(r[3]).strip() if r[3] is not None else ""
    nom = re.sub(r"\s+", " ", nom)
    prenom = re.sub(r"\s+", " ", prenom)
    if nom and prenom:
        students.append({"nom": nom, "prenom": prenom})
print(f"xlsx: {len(students)} students parsed")

# --- 2) group row -----------------------------------------------------------
existing = rest("GET", "groups", params="?level=eq.5&rotation=eq.rot2")
if isinstance(existing, list) and existing:
    print("group 5/rot2 already present:", existing[0]["label"], existing[0]["chat_id"])
else:
    ins = rest("POST", "groups", {"level": LEVEL, "rotation": ROTATION, "label": LABEL, "chat_id": CHAT_ID})
    print("inserted group:", ins if not isinstance(ins, dict) or "__http_error__" not in ins else ins)
    if isinstance(ins, dict) and "__http_error__" in ins:
        sys.exit("group insert failed")

# --- 3) overlap check against existing students ------------------------------
def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


existing_all = rest("GET", "students", params="?select=nom,prenom,level,rotation&limit=100000")
if not isinstance(existing_all, list):
    sys.exit(f"cannot list students: {existing_all}")
existing_keys = {norm(f"{s['nom']} {s['prenom']}") for s in existing_all}
rot2_keys = {norm(f"{s['nom']} {s['prenom']}") for s in students}
overlap = sorted(existing_keys & rot2_keys)
if overlap:
    print(f"WARNING: {len(overlap)} name(s) already in DB (any rotation):")
    for o in overlap:
        print("   ", o)
# ignore joined rot1 vs new — rotations are disjoint cohorts; report only

# --- 4) students (skip rows already present in rot2) --------------------------
existing_rot2 = {
    norm(f"{s['nom']} {s['prenom']}")
    for s in existing_all
    if s.get("level") == LEVEL and s.get("rotation") == ROTATION
}
to_insert = [s for s in students if norm(f"{s['nom']} {s['prenom']}") not in existing_rot2]
print(f"to insert: {len(to_insert)} (rot2 already has {len(existing_rot2)})")

BATCH = 200
for i in range(0, len(to_insert), BATCH):
    chunk = to_insert[i : i + BATCH]
    payload = [
        {"nom": s["nom"], "prenom": s["prenom"], "level": LEVEL, "rotation": ROTATION, "joined": False}
        for s in chunk
    ]
    res = rest("POST", "students", payload)
    if isinstance(res, dict) and "__http_error__" in res:
        sys.exit(f"insert batch {i // BATCH} failed: {res}")
    print(f"  inserted batch {i // BATCH + 1}: {len(chunk)}")

# --- 5) verify ---------------------------------------------------------------
check = rest("GET", "students", params="?select=nom,prenom&level=eq.5&rotation=eq.rot2")
print("verify: students in DB 5/rot2 =", len(check) if isinstance(check, list) else check)
groups = rest("GET", "groups", params="?select=level,rotation,label,chat_id&order=level.desc")
if isinstance(groups, list):
    for g in groups:
        print("verify group:", g)